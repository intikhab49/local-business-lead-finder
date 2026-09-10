import asyncio
import csv
import io
import sys
import threading
from pathlib import Path
from unittest.mock import patch

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.database import Base
from app.monitoring.metrics import get_metrics, reset_metrics

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"


@pytest.fixture()
def db_session(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'leads.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def api_client(tmp_path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.routes.export import router as export_router
    from app.api.routes.history import router as history_router
    from app.api.routes.leads import router
    from app.db.database import get_db

    engine = create_engine(
        f"sqlite:///{tmp_path / 'api_leads.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)

    def override_get_db():
        session = SessionLocal()
        try:
            yield session
        finally:
            session.close()

    app = FastAPI()
    app.include_router(router)
    app.include_router(export_router)
    app.include_router(history_router)
    app.dependency_overrides[get_db] = override_get_db

    with TestClient(app) as client:
        yield client
    engine.dispose()


def _lead_data(**overrides):
    data = {
        "fsq_place_id": "fsq_abc",
        "business_name": "Acme Cafe",
        "address": "123 Main St, Springfield",
        "phone": "555-0100",
        "website": "https://acme.example.com",
        "email": "hello@acme.example.com",
        "linkedin": "https://linkedin.com/company/acme",
        "facebook": "https://facebook.com/acme",
        "instagram": "https://instagram.com/acme",
        "twitter": "https://x.com/acme",
        "youtube": "https://youtube.com/@acme",
        "confidence_score": 0.9,
        "source": "foursquare",
        "review_status": "approved",
        "reviewer": "Alice",
    }
    data.update(overrides)
    return data


class TestLeadRepository:
    def test_create_and_get(self, db_session):
        from app.db import repository

        lead = repository.create_lead(db_session, _lead_data())
        assert lead.id is not None
        fetched = repository.get_lead(db_session, lead.id)
        assert fetched is not None
        assert fetched.business_name == "Acme Cafe"
        assert fetched.review_status == "approved"

    def test_get_missing_returns_none(self, db_session):
        from app.db import repository

        assert repository.get_lead(db_session, 999) is None

    def test_get_by_fsq(self, db_session):
        from app.db import repository

        repository.create_lead(db_session, _lead_data(fsq_place_id="uniq_1"))
        found = repository.get_lead_by_fsq(db_session, "uniq_1")
        assert found is not None
        assert repository.get_lead_by_fsq(db_session, "nope") is None

    def test_list_order_and_limits(self, db_session):
        from app.db import repository

        for i in range(5):
            repository.create_lead(db_session, _lead_data(fsq_place_id=f"fsq_{i}"))
        leads = repository.list_leads(db_session, limit=2, offset=0)
        assert len(leads) == 2
        assert leads[0].fsq_place_id == "fsq_4"
        assert repository.count_leads(db_session) == 5

    def test_list_search(self, db_session):
        from app.db import repository

        repository.create_lead(db_session, _lead_data(business_name="Pizza Palace"))
        repository.create_lead(db_session, _lead_data(business_name="Taco Town", fsq_place_id="fsq_2"))
        results = repository.list_leads(db_session, search="pizza")
        assert len(results) == 1
        assert results[0].business_name == "Pizza Palace"

    def test_list_review_status_filter(self, db_session):
        from app.db import repository

        repository.create_lead(db_session, _lead_data(review_status="approved"))
        repository.create_lead(db_session, _lead_data(review_status="edited", fsq_place_id="fsq_2"))
        results = repository.list_leads(db_session, review_status="edited")
        assert len(results) == 1
        assert results[0].review_status == "edited"

    def test_update_lead(self, db_session):
        from app.db import repository

        lead = repository.create_lead(db_session, _lead_data())
        updated = repository.update_lead(db_session, lead, {"email": "new@acme.example.com"})
        assert updated.email == "new@acme.example.com"
        assert repository.get_lead(db_session, lead.id).email == "new@acme.example.com"

    def test_update_ignores_unknown_fields(self, db_session):
        from app.db import repository

        lead = repository.create_lead(db_session, _lead_data())
        updated = repository.update_lead(db_session, lead, {"not_a_column": 42})
        assert updated.business_name == "Acme Cafe"

    def test_delete_lead(self, db_session):
        from app.db import repository

        lead = repository.create_lead(db_session, _lead_data())
        repository.delete_lead(db_session, lead)
        assert repository.get_lead(db_session, lead.id) is None
        assert repository.count_leads(db_session) == 0

    def test_stats(self, db_session):
        from app.db import repository

        repository.create_lead(db_session, _lead_data())
        repository.create_lead(db_session, _lead_data(fsq_place_id="fsq_2", email=None, website=None))
        stats = repository.get_stats(db_session)
        assert stats["total"] == 2
        assert stats["by_status"].get("approved") == 2
        assert stats["with_email"] == 1
        assert stats["with_website"] == 1
        assert stats["avg_confidence"] == 0.9


class TestLeadPersistenceService:
    def test_clean_categories_removes_generic_labels(self):
        from app.services.lead_persistence_service import clean_categories

        assert clean_categories("restaurant, point_of_interest, coffee_shop, establishment") == (
            "restaurant, coffee_shop"
        )

    def test_clean_categories_handles_edge_cases(self):
        from app.services.lead_persistence_service import clean_categories

        assert clean_categories(None) == ""
        assert clean_categories("") == ""
        assert clean_categories("point_of_interest") == ""
        assert clean_categories("establishment") == ""
        assert clean_categories("cafe") == "cafe"
        assert clean_categories("Establishment, Point_of_Interest, Cafe") == "Cafe"

    def test_clean_categories_list_removes_generic_labels(self):
        from app.services.lead_persistence_service import clean_categories_list

        assert clean_categories_list(
            ["educational_institution", "point_of_interest", "establishment"]
        ) == ["educational_institution"]
        assert clean_categories_list(["cafe", "point_of_interest"]) == ["cafe"]
        assert clean_categories_list([]) == []
        assert clean_categories_list(None) == []

    def test_clean_categories_keeps_database_record_unchanged(self, db_session):
        from app.db import repository
        from app.services.lead_persistence_service import clean_categories

        repository.create_lead(
            db_session,
            _lead_data(fsq_place_id="fsq_cat", categories="restaurant, point_of_interest"),
        )
        lead = repository.get_lead_by_fsq(db_session, "fsq_cat")
        assert lead.categories == "restaurant, point_of_interest"
        assert clean_categories(lead.categories) == "restaurant"

    def test_save_creates_lead(self, db_session):
        reset_metrics()
        from app.services.lead_persistence_service import save_reviewed_lead

        lead = save_reviewed_lead(
            db_session,
            fsq_place_id="fsq_save",
            lead_data={
                "business_name": "Saved Cafe",
                "address": "9 Elm St",
                "email": "c@c.com",
                "confidence_score": 0.85,
                "website": "https://c.example.com",
            },
            review_status="approved",
            reviewer="Bob",
        )
        assert lead.id is not None
        assert lead.review_status == "approved"
        assert lead.business_name == "Saved Cafe"
        assert get_metrics().get_counter("leads_created") == 1

    def test_save_updates_existing(self, db_session):
        reset_metrics()
        from app.services.lead_persistence_service import save_reviewed_lead

        save_reviewed_lead(
            db_session,
            fsq_place_id="fsq_save",
            lead_data={"business_name": "First Name"},
            review_status="approved",
        )
        updated = save_reviewed_lead(
            db_session,
            fsq_place_id="fsq_save",
            lead_data={"business_name": "Renamed Cafe"},
            review_status="edited",
        )
        assert updated.business_name == "Renamed Cafe"
        assert updated.review_status == "edited"
        assert get_metrics().get_counter("leads_updated") == 1

    def test_save_is_thread_safe(self, tmp_path):
        from app.services.lead_persistence_service import save_reviewed_lead

        engine = create_engine(
            f"sqlite:///{tmp_path / 'threads.db'}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(bind=engine)
        SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)

        results = []
        errors = []

        def worker(index: int):
            try:
                session = SessionLocal()
                try:
                    lead = save_reviewed_lead(
                        session,
                        fsq_place_id=f"fsq_t{index}",
                        lead_data={"business_name": f"Thread Cafe {index}"},
                        review_status="approved",
                    )
                    results.append(lead.id)
                finally:
                    session.close()
            except Exception as exc:  # noqa: BLE001
                errors.append(exc)

        threads = [threading.Thread(target=worker, args=(i,)) for i in range(4)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join()

        assert not errors
        assert len(results) == 4
        assert len(set(results)) == 4


class TestLeadsEndpoints:
    def test_create_lead(self, api_client):
        response = api_client.post("/leads", json=_lead_data())
        assert response.status_code == 201
        data = response.json()
        assert data["id"] is not None
        assert data["business_name"] == "Acme Cafe"
        assert data["review_status"] == "approved"

    def test_create_lead_idempotent_on_fsq(self, api_client):
        first = api_client.post("/leads", json=_lead_data())
        assert first.status_code == 201
        second = api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_abc"))
        assert second.status_code == 201
        assert second.json()["id"] == first.json()["id"]

    def test_list_leads_empty(self, api_client):
        response = api_client.get("/leads")
        assert response.status_code == 200
        data = response.json()
        assert data["total"] == 0
        assert data["items"] == []

    def test_list_leads_returns_cleaned_categories(self, api_client):
        api_client.post(
            "/leads",
            json=_lead_data(
                fsq_place_id="fsq_clean",
                categories="restaurant, point_of_interest, coffee_shop, establishment",
            ),
        )
        response = api_client.get("/leads")
        assert response.status_code == 200
        items = response.json()["items"]
        assert len(items) == 1
        assert items[0]["categories"] == "restaurant, coffee_shop"

    def test_get_lead_returns_cleaned_categories(self, api_client):
        created = api_client.post(
            "/leads",
            json=_lead_data(
                fsq_place_id="fsq_clean",
                categories="point_of_interest, establishment",
            ),
        ).json()
        response = api_client.get(f"/leads/{created['id']}")
        assert response.status_code == 200
        assert response.json()["categories"] == ""

    def test_list_leads_pagination(self, api_client):
        for i in range(3):
            api_client.post("/leads", json=_lead_data(fsq_place_id=f"fsq_{i}"))
        response = api_client.get("/leads", params={"limit": 2, "offset": 0})
        data = response.json()
        assert data["total"] == 3
        assert len(data["items"]) == 2
        assert data["items"][0]["fsq_place_id"] == "fsq_2"

    def test_list_leads_search(self, api_client):
        api_client.post("/leads", json=_lead_data(business_name="Pizza Palace"))
        api_client.post("/leads", json=_lead_data(business_name="Taco Town", fsq_place_id="fsq_2"))
        response = api_client.get("/leads", params={"search": "pizza"})
        data = response.json()
        assert data["total"] == 1
        assert data["items"][0]["business_name"] == "Pizza Palace"

    def test_get_lead(self, api_client):
        created = api_client.post("/leads", json=_lead_data()).json()
        response = api_client.get(f"/leads/{created['id']}")
        assert response.status_code == 200
        assert response.json()["business_name"] == "Acme Cafe"

    def test_get_lead_404(self, api_client):
        response = api_client.get("/leads/9999")
        assert response.status_code == 404
        assert response.json()["detail"]["error"] == "NotFound"

    def test_update_lead(self, api_client):
        created = api_client.post("/leads", json=_lead_data()).json()
        response = api_client.put(
            f"/leads/{created['id']}", json={"email": "new@acme.example.com"}
        )
        assert response.status_code == 200
        assert response.json()["email"] == "new@acme.example.com"

    def test_delete_lead(self, api_client):
        created = api_client.post("/leads", json=_lead_data()).json()
        response = api_client.delete(f"/leads/{created['id']}")
        assert response.status_code == 204
        assert api_client.get(f"/leads/{created['id']}").status_code == 404

    def test_delete_leads_respects_filters(self, api_client):
        api_client.post(
            "/leads", json=_lead_data(fsq_place_id="pizza", business_name="Pizza Palace")
        )
        api_client.post(
            "/leads", json=_lead_data(fsq_place_id="taco", business_name="Taco Town")
        )

        response = api_client.delete("/leads", params={"search": "pizza"})

        assert response.status_code == 200
        assert response.json() == {"deleted": 1}
        remaining = api_client.get("/leads", params={"limit": 10}).json()["items"]
        assert [lead["business_name"] for lead in remaining] == ["Taco Town"]

    def test_lead_stats(self, api_client):
        api_client.post("/leads", json=_lead_data())
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_2", email=None))
        response = api_client.get("/leads/stats")
        assert response.status_code == 200
        data = response.json()
        assert data["total"] == 2
        assert data["by_status"]["approved"] == 2
        assert data["with_email"] == 1

    def test_invalid_create_rejected(self, api_client):
        response = api_client.post("/leads", json={"confidence_score": 5.0})
        assert response.status_code == 422


class TestLeadsBatchEndpoint:
    def test_batch_creates_all_leads(self, api_client):
        items = [_lead_data(fsq_place_id=f"fsq_{i}", business_name=f"Biz {i}") for i in range(3)]
        response = api_client.post("/leads/batch", json={"items": items})
        assert response.status_code == 200
        result = response.json()
        assert result["created"] == 3
        assert result["updated"] == 0
        assert result["failed"] == 0
        assert result["run_id"] is not None
        listing = api_client.get("/leads", params={"limit": 10})
        assert listing.json()["total"] == 3

    def test_batch_dedupes_updates_existing(self, api_client):
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_same", business_name="Original"))
        response = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_same", business_name="Renamed"), _lead_data(fsq_place_id="fsq_new")]},
        )
        result = response.json()
        assert result["created"] == 1
        assert result["updated"] == 1
        listing = api_client.get("/leads", params={"limit": 10})
        assert listing.json()["total"] == 2
        by_fsq = {item["fsq_place_id"]: item for item in listing.json()["items"]}
        assert by_fsq["fsq_same"]["business_name"] == "Renamed"

    def test_batch_skips_missing_fsq(self, api_client):
        response = api_client.post(
            "/leads/batch",
            json={"items": [{"business_name": "No fsq"}, _lead_data(fsq_place_id="fsq_1")]},
        )
        result = response.json()
        assert result["failed"] == 1
        assert result["created"] == 1

    def test_batch_empty(self, api_client):
        response = api_client.post("/leads/batch", json={"items": []})
        assert response.status_code == 200
        assert response.json() == {"created": 0, "updated": 0, "failed": 0, "run_id": None}

    def test_batch_stores_metadata_and_run_id(self, api_client):
        items = [_lead_data(fsq_place_id=f"fsq_{i}", business_name=f"Biz {i}") for i in range(2)]
        response = api_client.post(
            "/leads/batch",
            json={"items": items, "location": "Austin, TX", "business_category": "coffee"},
        )
        result = response.json()
        assert result["created"] == 2
        assert result["run_id"] is not None
        listing = api_client.get("/leads", params={"run_id": result["run_id"], "limit": 10})
        data = listing.json()
        assert data["total"] == 2
        assert all(item["run_id"] == result["run_id"] for item in data["items"])

    def test_leads_batches_lists_save_all_runs(self, api_client):
        api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_1")], "location": "Austin, TX", "business_category": "coffee"},
        )
        api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_2", business_name="Taco Town")], "location": "Denver, CO", "business_category": "food"},
        )
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_individual", business_name="Single"))

        response = api_client.get("/leads/batches")
        assert response.status_code == 200
        data = response.json()
        assert data["total"] == 2
        assert data["ungrouped_total"] == 1
        assert {batch["location"] for batch in data["items"]} == {"Austin, TX", "Denver, CO"}
        assert {batch["business_category"] for batch in data["items"]} == {"coffee", "food"}
        assert all(batch["lead_count"] == 1 for batch in data["items"])

    def test_ungrouped_filter_excludes_batch_leads(self, api_client):
        batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_batch")], "location": "Austin, TX"},
        ).json()
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_individual", business_name="Single"))

        listing = api_client.get("/leads", params={"ungrouped": True, "limit": 10})
        data = listing.json()
        assert data["total"] == 1
        assert data["items"][0]["fsq_place_id"] == "fsq_individual"
        assert data["items"][0]["run_id"] is None

        batch_listing = api_client.get("/leads", params={"run_id": batch["run_id"], "limit": 10})
        assert batch_listing.json()["total"] == 1
        assert batch_listing.json()["items"][0]["fsq_place_id"] == "fsq_batch"

    def test_batches_empty(self, api_client):
        response = api_client.get("/leads/batches")
        assert response.status_code == 200
        data = response.json()
        assert data["items"] == []
        assert data["total"] == 0
        assert data["ungrouped_total"] == 0


class TestSavedLeadsRegression:
    """Batched leads must remain visible after run history is cleared.

    Leads saved via /leads/batch keep their run_id. If the run_history rows
    are later cleared (e.g. the "Clear history" action), the leads become
    orphaned: they no longer show as a batch, so they must still be returned
    as individual/ungrouped leads on the Saved Leads page.
    """

    def test_orphaned_batch_leads_survive_run_history_clear(self, api_client):
        batch = api_client.post(
            "/leads/batch",
            json={
                "items": [
                    _lead_data(fsq_place_id="fsq_1", business_name="Biz One"),
                    _lead_data(fsq_place_id="fsq_2", business_name="Biz Two"),
                ],
                "location": "Austin, TX",
                "business_category": "coffee",
            },
        ).json()
        assert batch["created"] == 2
        assert batch["run_id"] is not None

        response = api_client.delete("/history")
        assert response.status_code == 200
        assert response.json() == {"deleted": 1}

        batches = api_client.get("/leads/batches").json()
        assert batches["items"] == []
        assert batches["total"] == 0
        assert batches["ungrouped_total"] == 2

        ungrouped = api_client.get("/leads", params={"ungrouped": True, "limit": 10}).json()
        assert ungrouped["total"] == 2
        names = {item["business_name"] for item in ungrouped["items"]}
        assert names == {"Biz One", "Biz Two"}
        assert all(item["run_id"] is not None for item in ungrouped["items"])


class TestExportEndpoints:
    def test_export_csv(self, api_client):
        api_client.post("/leads", json=_lead_data())
        response = api_client.get("/export/csv")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert "attachment; filename=leads_" in response.headers["content-disposition"]
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["business_name"] == "Acme Cafe"
        assert rows[0]["email"] == "hello@acme.example.com"

    def test_export_csv_empty(self, api_client):
        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert rows == []

    def test_export_excel(self, api_client):
        from openpyxl import load_workbook

        api_client.post("/leads", json=_lead_data())
        response = api_client.get("/export/excel")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        assert worksheet.title == "Leads"
        headers = [cell.value for cell in worksheet[1]]
        assert "business_name" in headers
        assert worksheet.max_row == 2
        row = {headers[i]: worksheet[2][i].value for i in range(len(headers))}
        assert row["business_name"] == "Acme Cafe"

    def test_export_excel_empty(self, api_client):
        from openpyxl import load_workbook

        response = api_client.get("/export/excel")
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.active.max_row == 1

    def test_export_csv_contains_all_leads(self, api_client):
        for i in range(5):
            api_client.post("/leads", json=_lead_data(fsq_place_id=f"fsq_{i}", business_name=f"Biz {i}"))
        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 5
        names = {row["business_name"] for row in rows}
        assert names == {f"Biz {i}" for i in range(5)}

    def test_export_csv_includes_new_columns(self, api_client):
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_cols"))
        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        row = rows[0]
        assert row["business_name"] == "Acme Cafe"
        assert "city" in row
        assert "state" in row
        assert "country" in row
        assert "categories" in row
        assert "linkedin" in row
        assert "facebook" in row
        assert "instagram" in row
        assert "twitter" in row
        assert "youtube" in row
        assert row["linkedin"] == "https://linkedin.com/company/acme"
        assert row["youtube"] == "https://youtube.com/@acme"
        assert "created_at" in row
        assert "updated_at" in row

    def test_export_excel_contains_all_leads(self, api_client):
        from openpyxl import load_workbook

        for i in range(4):
            api_client.post("/leads", json=_lead_data(fsq_place_id=f"fsq_{i}", business_name=f"Biz {i}"))
        response = api_client.get("/export/excel")
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        assert worksheet.max_row == 5
        names = {worksheet.cell(row=r, column=3).value for r in range(2, 6)}
        assert names == {f"Biz {i}" for i in range(4)}

    def test_export_csv_ids_sequential_from_one(self, api_client):
        # Create leads with non-contiguous DB primary keys to prove the exported
        # id is sequential and never the DB id.
        first = api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_a", business_name="A")).json()
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_b", business_name="B")).json()
        api_client.delete(f"/leads/{first['id']}")
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_c", business_name="C")).json()

        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert [row["id"] for row in rows] == ["1", "2"]
        names = {row["business_name"] for row in rows}
        assert names == {"B", "C"}

    def test_export_excel_ids_sequential_from_one(self, api_client):
        from openpyxl import load_workbook

        first = api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_a", business_name="A")).json()
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_b", business_name="B")).json()
        api_client.delete(f"/leads/{first['id']}")
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_c", business_name="C")).json()

        response = api_client.get("/export/excel")
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        id_col = headers.index("id") + 1
        ids = [worksheet.cell(row=r, column=id_col).value for r in range(2, worksheet.max_row + 1)]
        assert ids == [1, 2]

    def test_export_csv_ids_start_at_one_per_batch(self, api_client):
        first_batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b1", business_name="B1"), _lead_data(fsq_place_id="fsq_b2", business_name="B2")], "location": "Austin, TX"},
        ).json()
        second_batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b3", business_name="B3")], "location": "Denver, CO"},
        ).json()

        first_export = api_client.get("/export/csv", params={"run_id": first_batch["run_id"]})
        first_rows = list(csv.DictReader(io.StringIO(first_export.text)))
        assert [row["id"] for row in first_rows] == ["1", "2"]

        second_export = api_client.get("/export/csv", params={"run_id": second_batch["run_id"]})
        second_rows = list(csv.DictReader(io.StringIO(second_export.text)))
        assert [row["id"] for row in second_rows] == ["1"]

    def test_export_excel_ids_start_at_one_per_batch(self, api_client):
        from openpyxl import load_workbook

        first_batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b1", business_name="B1"), _lead_data(fsq_place_id="fsq_b2", business_name="B2")], "location": "Austin, TX"},
        ).json()
        second_batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b3", business_name="B3")], "location": "Denver, CO"},
        ).json()

        def export_ids(run_id):
            response = api_client.get("/export/excel", params={"run_id": run_id})
            workbook = load_workbook(io.BytesIO(response.content))
            worksheet = workbook.active
            headers = [cell.value for cell in worksheet[1]]
            id_col = headers.index("id") + 1
            return [worksheet.cell(row=r, column=id_col).value for r in range(2, worksheet.max_row + 1)]

        assert export_ids(first_batch["run_id"]) == [1, 2]
        assert export_ids(second_batch["run_id"]) == [1]

    def test_export_csv_filters_generic_categories(self, api_client):
        api_client.post(
            "/leads",
            json=_lead_data(
                fsq_place_id="fsq_cat",
                categories="restaurant, point_of_interest, coffee_shop, establishment",
            ),
        )
        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["categories"] == "restaurant, coffee_shop"

    def test_export_excel_filters_generic_categories(self, api_client):
        from openpyxl import load_workbook

        api_client.post(
            "/leads",
            json=_lead_data(
                fsq_place_id="fsq_cat",
                categories="restaurant, point_of_interest, coffee_shop, establishment",
            ),
        )
        response = api_client.get("/export/excel")
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        headers = [cell.value for cell in worksheet[1]]
        cat_col = headers.index("categories") + 1
        assert worksheet.cell(row=2, column=cat_col).value == "restaurant, coffee_shop"

    def test_export_keeps_database_id_untouched(self, api_client):
        created = api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_id", categories="bar")).json()
        response = api_client.get("/export/csv")
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert rows[0]["id"] == "1"
        fetched = api_client.get(f"/leads/{created['id']}").json()
        assert fetched["id"] == created["id"]

    def test_export_respects_filters(self, api_client):
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_a", business_name="Approved Cafe", review_status="approved"))
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_b", business_name="Edited Cafe", review_status="edited"))
        response = api_client.get("/export/csv", params={"review_status": "approved"})
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["business_name"] == "Approved Cafe"

    def test_export_csv_respects_run_id(self, api_client):
        batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b1", business_name="Batch One")], "location": "Austin, TX"},
        ).json()
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_ind", business_name="Individual"))

        response = api_client.get("/export/csv", params={"run_id": batch["run_id"]})
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["business_name"] == "Batch One"

    def test_export_csv_ungrouped_excludes_batch(self, api_client):
        api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b1", business_name="Batch One")], "location": "Austin, TX"},
        )
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_ind", business_name="Individual"))

        response = api_client.get("/export/csv", params={"ungrouped": True})
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["business_name"] == "Individual"

    def test_export_excel_respects_run_id(self, api_client):
        from openpyxl import load_workbook

        batch = api_client.post(
            "/leads/batch",
            json={"items": [_lead_data(fsq_place_id="fsq_b1", business_name="Batch One")], "location": "Austin, TX"},
        ).json()
        api_client.post("/leads", json=_lead_data(fsq_place_id="fsq_ind", business_name="Individual"))

        response = api_client.get("/export/excel", params={"run_id": batch["run_id"]})
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        assert worksheet.max_row == 2
        headers = [cell.value for cell in worksheet[1]]
        row = {headers[i]: worksheet[2][i].value for i in range(len(headers))}
        assert row["business_name"] == "Batch One"


class TestExportMetrics:
    def test_export_counters_incremented(self, api_client):
        reset_metrics()
        api_client.post("/leads", json=_lead_data())
        api_client.get("/export/csv")
        api_client.get("/export/excel")
        metrics = get_metrics()
        assert metrics.get_counter("leads_exported_csv") == 1
        assert metrics.get_counter("leads_exported_excel") == 1


class TestReviewAutoSaveIntegration:
    def test_approve_persists_lead(self, tmp_path):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from sqlalchemy import create_engine
        from sqlalchemy.orm import sessionmaker

        from app.api.routes.review import review_agent, router
        from app.db import repository
        from app.db.database import get_db

        review_agent.clear()
        try:
            engine = create_engine(
                f"sqlite:///{tmp_path / 'review.db'}",
                connect_args={"check_same_thread": False},
            )
            Base.metadata.create_all(bind=engine)
            SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)

            def override_get_db():
                session = SessionLocal()
                try:
                    yield session
                finally:
                    session.close()

            app = FastAPI()
            app.include_router(router)
            app.dependency_overrides[get_db] = override_get_db

            with TestClient(app) as client:
                asyncio.run(
                    review_agent.submit(
                        "fsq_review1",
                        reviewer="system",
                        enriched={
                            "business_name": "Review Cafe",
                            "address": "1 Review St",
                            "email": "review@cafe.com",
                            "confidence_score": 0.4,
                            "source": "foursquare",
                        },
                    )
                )
                response = client.post(
                    "/review/fsq_review1/approve", json={"reviewer": "Alice"}
                )
                assert response.status_code == 200
                assert response.json()["status"] == "approved"

            session = SessionLocal()
            try:
                lead = repository.get_lead_by_fsq(session, "fsq_review1")
                assert lead is not None
                assert lead.review_status == "approved"
                assert lead.business_name == "Review Cafe"
                assert lead.reviewer == "Alice"
            finally:
                session.close()
                engine.dispose()
        finally:
            review_agent.clear()


class TestSavedLeadsPageUI:
    @staticmethod
    def _load_frontend_app():
        import importlib.util

        if "frontend_app" in sys.modules:
            return sys.modules["frontend_app"]
        sys.path.insert(0, str(FRONTEND_DIR))
        try:
            spec = importlib.util.spec_from_file_location(
                "frontend_app", FRONTEND_DIR / "app.py"
            )
            module = importlib.util.module_from_spec(spec)
            sys.modules["frontend_app"] = module
            spec.loader.exec_module(module)
        finally:
            sys.path.remove(str(FRONTEND_DIR))
        return module

    @staticmethod
    def _fixture_lead(lead_id=1, name="Acme Cafe"):
        return {
            "id": lead_id,
            "fsq_place_id": f"fsq_{lead_id}",
            "business_name": name,
            "address": "123 Main St, Springfield",
            "phone": "555-0100",
            "website": "https://acme.example.com",
            "email": "hello@acme.example.com",
            "confidence_score": 0.9,
            "source": "foursquare",
            "review_status": "approved",
            "run_id": None,
        }

    @staticmethod
    def _fixture_batch(run_id=1, location="Austin, TX", category="coffee", lead_count=1):
        return {
            "run_id": run_id,
            "location": location,
            "business_category": category,
            "status": "success",
            "total_results": lead_count,
            "lead_count": lead_count,
            "started_at": "2026-08-11T10:00:00",
            "finished_at": "2026-08-11T10:00:05",
        }

    @staticmethod
    def _fixture_stats():
        return {
            "total": 1,
            "by_status": {"approved": 1},
            "avg_confidence": 0.9,
            "with_email": 1,
            "with_website": 1,
        }

    @staticmethod
    def _page(**overrides):
        from contextlib import contextmanager

        from streamlit.testing.v1 import AppTest

        harness = str(Path(__file__).parent / "_leads_harness.py")

        @contextmanager
        def _run():
            appmod = TestSavedLeadsPageUI._load_frontend_app()
            with patch.multiple(appmod, **overrides):
                at = AppTest.from_file(harness, default_timeout=20)
                at.run()
                yield at

        return _run()

    def test_empty_state(self):
        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 0},
            list_leads=lambda **kwargs: {"items": [], "total": 0, "limit": 50, "offset": 0},
            export_csv=lambda **kwargs: ("leads.csv", "id,business_name\n"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"bytes"),
        ) as at:
            assert not at.exception
            assert at.info
            assert len(at.get("download_button")) == 0

    def test_orphaned_leads_rendered_not_empty_state(self):
        """Batched leads whose run history was cleared must still show up.

        Regression: after clearing run history, /leads/batches returns no
        batches but /leads?ungrouped=true returns the orphaned leads. The
        page must render them instead of the "No saved lead yet" empty state.
        """

        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped"):
                return {
                    "items": [
                        self._fixture_lead(lead_id=1, name="Biz One"),
                        self._fixture_lead(lead_id=2, name="Biz Two"),
                    ],
                    "total": 2,
                    "limit": 200,
                    "offset": 0,
                }
            return {"items": [], "total": 0, "limit": 200, "offset": 0}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 2},
            list_leads=fake_list_leads,
            export_csv=lambda **kwargs: ("leads.csv", "id,business_name\n"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"bytes"),
        ) as at:
            assert not at.exception
            assert not at.info
            assert at.expander
            assert at.dataframe
            df = at.dataframe[0].value
            assert set(df["Business Name"].tolist()) >= {"Biz One", "Biz Two"}

    def test_renders_batched_leads(self):
        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped") or kwargs.get("run_id") is None:
                return {"items": [], "total": 0, "limit": 200, "offset": 0}
            return {
                "items": [self._fixture_lead(lead_id=1)],
                "total": 1,
                "limit": 200,
                "offset": 0,
            }

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {
                "items": [self._fixture_batch(run_id=1, location="Austin, TX", category="coffee")],
                "total": 1,
                "ungrouped_total": 0,
            },
            list_leads=fake_list_leads,
            export_csv=lambda **kwargs: ("leads.csv", "id,business_name\n1,Acme Cafe\n"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"bytes"),
        ) as at:
            assert not at.exception
            assert at.expander
            assert len(at.get("download_button")) == 2

    def test_cleaned_categories_displayed(self):
        cleaned_lead = self._fixture_lead(lead_id=1)
        cleaned_lead["categories"] = "restaurant, coffee_shop"

        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped"):
                return {"items": [cleaned_lead], "total": 1, "limit": 200, "offset": 0}
            return {"items": [], "total": 0, "limit": 200, "offset": 0}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 1},
            list_leads=fake_list_leads,
        ) as at:
            assert not at.exception
            assert at.expander
            assert at.dataframe
            df = at.dataframe[0].value
            assert df.loc[0, "Categories"] == "restaurant, coffee_shop"

    def test_batch_metadata_rendered(self):
        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped") or kwargs.get("run_id") is None:
                return {"items": [], "total": 0, "limit": 200, "offset": 0}
            return {
                "items": [self._fixture_lead(lead_id=1)],
                "total": 1,
                "limit": 200,                "offset": 0,
            }

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {
                "items": [self._fixture_batch(run_id=1, location="Austin, TX", category="coffee")],
                "total": 1,
                "ungrouped_total": 0,
            },
            list_leads=fake_list_leads,
        ) as at:
            assert not at.exception
        rendered = " ".join(str(md.value) for md in at.markdown)
        captions = " ".join(str(cap.value) for cap in at.caption)
        assert "Austin, TX" in rendered
        assert "coffee" in rendered
        assert "lead(s)" in captions

    def test_ungrouped_section_rendered(self):
        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped"):
                return {
                    "items": [self._fixture_lead(lead_id=1)],
                    "total": 1,
                    "limit": 200,
                    "offset": 0,
                }
            return {"items": [], "total": 0, "limit": 200, "offset": 0}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 1},
            list_leads=fake_list_leads,
            export_csv=lambda **kwargs: ("leads.csv", "id,business_name\n1,Acme Cafe\n"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"bytes"),
        ) as at:
            assert not at.exception
            assert at.expander
            assert len(at.get("download_button")) == 2
        rendered = " ".join(str(md.value) for md in at.markdown)
        assert "Individual Leads" in rendered

    def test_search_filters_passed(self):
        calls = {}

        def fake_list_leads(search="", review_status="", limit=200, offset=0, run_id=None, ungrouped=False):
            calls["search"] = search
            calls["review_status"] = review_status
            return {"items": [], "total": 0, "limit": limit, "offset": offset}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 0},
            list_leads=fake_list_leads,
        ) as at:
            at.text_input(key="leads_search").set_value("pizza")
            at.selectbox(key="leads_status").set_value("approved")
            at.run()
            assert not at.exception
        assert calls["search"] == "pizza"
        assert calls["review_status"] == "approved"

    def test_per_batch_export_uses_run_id(self):
        calls = []

        def fake_export_csv(**kwargs):
            calls.append(kwargs)
            return ("leads.csv", "id,business_name\n")

        def fake_export_excel(**kwargs):
            return ("leads.xlsx", b"bytes")

        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped") or kwargs.get("run_id") is None:
                return {"items": [], "total": 0, "limit": 200, "offset": 0}
            return {
                "items": [self._fixture_lead(lead_id=1)],
                "total": 1,
                "limit": 200,
                "offset": 0,
            }

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {
                "items": [self._fixture_batch(run_id=7, location="Austin, TX", category="coffee")],
                "total": 1,
                "ungrouped_total": 0,
            },
            list_leads=fake_list_leads,
            export_csv=fake_export_csv,
            export_excel=fake_export_excel,
        ) as at:
            assert not at.exception
            assert len(at.get("download_button")) == 2
        assert calls
        assert all(call["run_id"] == 7 for call in calls)
        assert all(not call["ungrouped"] for call in calls)

    def test_delete_flow(self):
        calls = {}

        def fake_delete(lead_id):
            calls["lead_id"] = lead_id

        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped"):
                return {
                    "items": [self._fixture_lead()],
                    "total": 1,
                    "limit": 200,
                    "offset": 0,
                }
            return {"items": [], "total": 0, "limit": 200, "offset": 0}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 1},
            list_leads=fake_list_leads,
            export_csv=lambda **kwargs: ("leads.csv", "data"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"data"),
            delete_lead=fake_delete,
        ) as at:
            at.button(key="delete_lead_1").click().run()
            assert not at.exception
        assert calls["lead_id"] == 1

    def test_delete_all_flow_uses_section_filters(self):
        calls = {}

        def fake_delete_leads(**kwargs):
            calls.update(kwargs)
            return 1

        def fake_list_leads(**kwargs):
            if kwargs.get("ungrouped"):
                return {
                    "items": [self._fixture_lead()],
                    "total": 1,
                    "limit": 200,
                    "offset": 0,
                }
            return {"items": [], "total": 0, "limit": 200, "offset": 0}

        with self._page(
            get_lead_stats=lambda: self._fixture_stats(),
            list_lead_batches=lambda: {"items": [], "total": 0, "ungrouped_total": 1},
            list_leads=fake_list_leads,
            export_csv=lambda **kwargs: ("leads.csv", "data"),
            export_excel=lambda **kwargs: ("leads.xlsx", b"data"),
            delete_leads=fake_delete_leads,
        ) as at:
            at.button(key="export_ungrouped_delete_all_btn").click().run()
            assert not at.exception

        assert calls == {
            "search": "",
            "review_status": "",
            "run_id": None,
            "ungrouped": True,
        }


class TestSearchPageSaveUI:
    @staticmethod
    def _page(**overrides):
        from contextlib import contextmanager

        from streamlit.testing.v1 import AppTest

        harness = str(Path(__file__).parent / "_search_harness.py")

        @contextmanager
        def _run():
            appmod = TestSavedLeadsPageUI._load_frontend_app()
            if overrides:
                with patch.multiple(appmod, **overrides):
                    at = AppTest.from_file(harness, default_timeout=20)
                    at.run()
                    yield at
            else:
                at = AppTest.from_file(harness, default_timeout=20)
                at.run()
                yield at

        return _run()

    @staticmethod
    def _fixture_business(fsq_id="fsq_1", name="Acme Cafe"):
        return {
            "fsq_place_id": fsq_id,
            "name": name,
            "formatted_address": "123 Main St, Springfield",
            "phone_number": "555-0100",
            "website": "https://acme.example.com",
            "categories": ["cafe"],
            "latitude": 40.0,
            "longitude": -74.0,
        }

    def test_initial_state(self):
        with self._page() as at:
            assert not at.exception
            assert at.info

    def test_search_results_render_save_buttons(self):
        with self._page(
            search_businesses=lambda **kwargs: {
                "businesses": [self._fixture_business()],
                "total_results": 1,
            },
            list_leads=lambda **kwargs: {
                "items": [],
                "total": 0,
                "limit": 200,
                "offset": 0,
            },
        ) as at:
            at.text_input(key="search_location").set_value("New York, NY")
            at.text_input(key="search_category").set_value("cafe")
            at.button(key="search_btn").click().run()
            assert not at.exception
            assert at.button(key="save_all_leads")
            assert at.button(key="save_lead_fsq_1")

    def test_email_shown_without_enrich(self):
        with self._page(
            search_businesses=lambda **kwargs: {
                "businesses": [self._fixture_business()],
                "total_results": 1,
            },
            list_leads=lambda **kwargs: {
                "items": [
                    {
                        "fsq_place_id": "fsq_1",
                        "email": "acme@example.com",
                    }
                ],
                "total": 1,
                "limit": 200,
                "offset": 0,
            },
        ) as at:
            at.text_input(key="search_location").set_value("New York, NY")
            at.text_input(key="search_category").set_value("cafe")
            at.button(key="search_btn").click().run()
            assert not at.exception
        rendered = " ".join(str(md.value) for md in at.markdown)
        assert "acme@example.com" in rendered

    def test_discover_more_button_removed(self):
        business = self._fixture_business()
        business.update(
            {
                "linkedin": "https://linkedin.com/company/acme",
                "facebook": "https://facebook.com/acme",
                "instagram": "https://instagram.com/acme",
                "twitter": "https://x.com/acme",
                "youtube": "https://youtube.com/@acme",
            }
        )
        with self._page(
            search_businesses=lambda **kwargs: {
                "businesses": [business],
                "total_results": 1,
            },
            list_leads=lambda **kwargs: {
                "items": [],
                "total": 0,
                "limit": 200,
                "offset": 0,
            },
        ) as at:
            at.text_input(key="search_location").set_value("New York, NY")
            at.text_input(key="search_category").set_value("cafe")
            at.button(key="search_btn").click().run()
            assert not at.exception
            assert not any(b.key.startswith("discover_more_") for b in at.button)
            assert any(b.key.startswith("save_lead_") for b in at.button)
        rendered = " ".join(str(md.value) for md in at.markdown)
        assert rendered.count("https://linkedin.com/company/acme") >= 1
        assert rendered.count("https://youtube.com/@acme") >= 1

    def test_save_single_lead(self):
        calls = []

        def fake_create(payload):
            calls.append(payload)
            return payload

        with self._page(
            search_businesses=lambda **kwargs: {
                "businesses": [self._fixture_business()],
                "total_results": 1,
            },
            list_leads=lambda **kwargs: {
                "items": [],
                "total": 0,
                "limit": 200,
                "offset": 0,
            },
            create_lead=fake_create,
        ) as at:
            at.text_input(key="search_location").set_value("New York, NY")
            at.text_input(key="search_category").set_value("cafe")
            at.button(key="search_btn").click().run()
            at.button(key="save_lead_fsq_1").click().run()
            assert not at.exception
        assert len(calls) == 1
        assert calls[0]["fsq_place_id"] == "fsq_1"
        assert calls[0]["business_name"] == "Acme Cafe"
        assert calls[0]["phone"] == "555-0100"

    def test_save_all_leads(self):
        calls = []

        def fake_save_batch(items, location="", business_category=""):
            calls.append({"items": items, "location": location, "business_category": business_category})
            return {"created": len(items), "updated": 0, "failed": 0, "run_id": 99}

        businesses = [
            self._fixture_business(fsq_id="fsq_1", name="Acme Cafe"),
            self._fixture_business(fsq_id="fsq_2", name="Taco Town"),
        ]
        with self._page(
            search_businesses=lambda **kwargs: {
                "businesses": businesses,
                "total_results": 2,
            },
            list_leads=lambda **kwargs: {
                "items": [],
                "total": 0,
                "limit": 200,
                "offset": 0,
            },
            save_leads_batch=fake_save_batch,
        ) as at:
            at.text_input(key="search_location").set_value("New York, NY")
            at.text_input(key="search_category").set_value("cafe")
            at.button(key="search_btn").click().run()
            at.button(key="save_all_leads").click().run()
            assert not at.exception
        assert len(calls) == 1
        call = calls[0]
        assert call["location"] == "New York, NY"
        assert call["business_category"] == "cafe"
        payloads = call["items"]
        assert {payload["fsq_place_id"] for payload in payloads} == {"fsq_1", "fsq_2"}
        assert {payload["source"] for payload in payloads} == {"google_places"}
