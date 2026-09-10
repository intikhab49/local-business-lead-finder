import csv
import io
from typing import Annotated

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.db import repository
from app.db.database import get_db
from app.db.models import Lead
from app.monitoring.metrics import get_metrics
from app.services.lead_persistence_service import clean_categories

router = APIRouter(prefix="/export", tags=["export"])

DbSession = Annotated[Session, Depends(get_db)]

_EXPORT_FIELDS = [
    "id",
    "fsq_place_id",
    "business_name",
    "address",
    "city",
    "state",
    "country",
    "phone",
    "website",
    "email",
    "linkedin",
    "facebook",
    "instagram",
    "twitter",
    "youtube",
    "categories",
    "latitude",
    "longitude",
    "confidence_score",
    "source",
    "review_status",
    "created_at",
    "updated_at",
]


def _rows(
    db: Session,
    search: str | None,
    review_status: str | None,
    run_id: int | None = None,
    ungrouped: bool = False,
) -> list[dict[str, object]]:
    leads = repository.list_leads(
        db,
        search=search,
        review_status=review_status,
        run_id=run_id,
        ungrouped=ungrouped,
        limit=None,
        offset=0,
    )
    return [_lead_row(lead, export_id=index + 1) for index, lead in enumerate(leads)]


def _lead_row(lead: Lead, export_id: int) -> dict[str, object]:
    """Serialize a Lead ORM object into an export-friendly dict.

    ``export_id`` is a sequential number starting at 1 for every export; the
    database primary key is never exposed.
    """
    created = lead.created_at.isoformat() if lead.created_at else None
    updated = lead.updated_at.isoformat() if lead.updated_at else None
    return {
        "id": export_id,
        "fsq_place_id": lead.fsq_place_id,
        "business_name": lead.business_name,
        "address": lead.address,
        "city": lead.city or "",
        "state": lead.state or "",
        "country": lead.country or "",
        "phone": lead.phone or "",
        "website": lead.website or "",
        "email": lead.email or "",
        "linkedin": lead.linkedin or "",
        "facebook": lead.facebook or "",
        "instagram": lead.instagram or "",
        "twitter": lead.twitter or "",
        "youtube": lead.youtube or "",
        "categories": clean_categories(lead.categories),
        "latitude": lead.latitude,
        "longitude": lead.longitude,
        "confidence_score": lead.confidence_score or 0.0,
        "source": lead.source or "",
        "review_status": lead.review_status or "",
        "created_at": created,
        "updated_at": updated,
    }


def _datetime_header() -> str:
    from datetime import UTC, datetime

    return datetime.now(UTC).strftime("%Y%m%d_%H%M%S")


@router.get("/csv", response_class=Response)
def export_csv(
    db: DbSession,
    search: str | None = Query(default=None, max_length=200),
    review_status: str | None = Query(default=None, max_length=50),
    run_id: int | None = Query(default=None, ge=1),
    ungrouped: bool = Query(default=False),
) -> Response:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=_EXPORT_FIELDS)
    writer.writeheader()
    for row in _rows(db, search, review_status, run_id, ungrouped):
        writer.writerow(row)

    filename = f"leads_{_datetime_header()}.csv"
    get_metrics().increment("leads_exported_csv")
    return Response(
        content=buffer.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/excel", response_class=Response)
def export_excel(
    db: DbSession,
    search: str | None = Query(default=None, max_length=200),
    review_status: str | None = Query(default=None, max_length=50),
    run_id: int | None = Query(default=None, ge=1),
    ungrouped: bool = Query(default=False),
) -> Response:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Leads"
    worksheet.append(_EXPORT_FIELDS)
    for row in _rows(db, search, review_status, run_id, ungrouped):
        worksheet.append([row[field] for field in _EXPORT_FIELDS])

    for column_index in range(1, len(_EXPORT_FIELDS) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 20

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = f"leads_{_datetime_header()}.xlsx"
    get_metrics().increment("leads_exported_excel")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
