import csv
import io
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.database import Base

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"


@pytest.fixture()
def db_session(tmp_path):
    engine = create_engine(
        f"sqlite:///{tmp_path / 'history.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture()
def api_client(tmp_path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.routes.history import router
    from app.db.database import get_db

    engine = create_engine(
        f"sqlite:///{tmp_path / 'api_history.db'}",
        connect_args={"check_same_thread": False},
    )
    Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)

    def override_get_db():
        session = SessionLocal()
        try:
            yield session
        finally:
            session.close()

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_db] = override_get_db

    def _seed(*items):
        session = SessionLocal()
        try:
            from app.db import repository

            created = []
            for item in items:
                created.append(repository.create_run_history(session, item))
            return created
        finally:
            session.close()

    with TestClient(app) as client:
        yield client, _seed
    engine.dispose()


def _run_data(**overrides):
    now = datetime.now(UTC)
    data = {
        "run_type": "search",
        "location": "New York, NY",
        "business_category": "restaurant",
        "status": "success",
        "total_tasks": 1,
        "total_results": 5,
        "processing_time": 1.25,
        "started_at": now - timedelta(seconds=2),
        "finished_at": now,
        "error_message": None,
    }
    data.update(overrides)
    return data


class TestRunHistoryRepository:
    def test_create_and_get(self, db_session):
        from app.db import repository

        run = repository.create_run_history(db_session, _run_data())
        assert run.id is not None
        fetched = repository.get_run_history(db_session, run.id)
        assert fetched is not None
        assert fetched.run_type == "search"
        assert fetched.status == "success"
        assert fetched.total_results == 5

    def test_get_missing_returns_none(self, db_session):
        from app.db import repository

        assert repository.get_run_history(db_session, 999) is None

    def test_finish_updates_fields(self, db_session):
        from app.db import repository

        run = repository.create_run_history(db_session, _run_data(status="running"))
        finished = repository.finish_run_history(
            db_session,
            run,
            status="failed",
            total_results=0,
            error_message="boom",
            processing_time=0.5,
        )
        assert finished.status == "failed"
        assert finished.error_message == "boom"
        assert finished.processing_time == 0.5
        assert finished.finished_at is not None

    def test_list_order_and_limits(self, db_session):
        from app.db import repository

        for i in range(5):
            repository.create_run_history(
                db_session, _run_data(location=f"City {i}", business_category="cafe")
            )
        runs = repository.list_run_history(db_session, limit=2, offset=0)
        assert len(runs) == 2
        assert runs[0].location == "City 4"
        assert repository.count_run_history(db_session) == 5

    def test_list_filters(self, db_session):
        from app.db import repository

        repository.create_run_history(
            db_session, _run_data(run_type="search", status="success", location="New York")
        )
        repository.create_run_history(
            db_session,
            _run_data(run_type="agent_search", status="failed", location="Austin"),
        )
        repository.create_run_history(
            db_session,
            _run_data(run_type="discover_more", status="success", location="Denver"),
        )
        assert repository.count_run_history(db_session, run_type="search") == 1
        assert repository.count_run_history(db_session, status="failed") == 1
        assert repository.count_run_history(db_session, search="austin") == 1
        assert repository.count_run_history(db_session, search="cafe") == 0
        assert repository.count_run_history(db_session, run_type="nope") == 0

    def test_stats(self, db_session):
        from app.db import repository

        repository.create_run_history(
            db_session, _run_data(status="success", processing_time=1.0)
        )
        repository.create_run_history(
            db_session, _run_data(status="success", processing_time=3.0)
        )
        repository.create_run_history(
            db_session, _run_data(status="failed", processing_time=2.0)
        )
        stats = repository.get_run_history_stats(db_session)
        assert stats["total_runs"] == 3
        assert stats["successful"] == 2
        assert stats["failed"] == 1
        assert stats["avg_duration"] == 2.0
        assert stats["by_status"]["success"] == 2

    def test_delete(self, db_session):
        from app.db import repository

        run = repository.create_run_history(db_session, _run_data())
        repository.delete_run_history(db_session, run)
        assert repository.get_run_history(db_session, run.id) is None
        assert repository.count_run_history(db_session) == 0

    def test_clear(self, db_session):
        from app.db import repository

        repository.create_run_history(db_session, _run_data())
        repository.create_run_history(db_session, _run_data())
        deleted = repository.clear_run_history(db_session)
        assert deleted == 2
        assert repository.count_run_history(db_session) == 0


class TestRunHistoryService:
    def test_start_and_finish(self, db_session):
        from app.db import repository
        from app.services.run_history_service import finish_run, start_run

        run = start_run(db_session, "search", location="NYC", business_category="gym")
        assert run.status == "running"
        assert run.location == "NYC"
        assert run.started_at is not None
        assert run.finished_at is None

        finished = finish_run(
            db_session,
            run,
            status="success",
            total_results=3,
            processing_time=0.4,
        )
        assert finished.status == "success"
        assert finished.total_results == 3
        assert finished.finished_at is not None
        assert repository.get_run_history(db_session, run.id).status == "success"

    def test_finish_truncates_error(self, db_session):
        from app.services.run_history_service import finish_run, start_run

        run = start_run(db_session, "discover_more", location="id-123")
        finished = finish_run(
            db_session,
            run,
            status="failed",
            error_message="x" * 5000,
        )
        assert len(finished.error_message) == 2000

    def test_mark_interrupted_runs(self, db_session):
        from app.db import repository
        from app.services.run_history_service import mark_interrupted_runs

        checkpoint = {"cells_planned": 36, "cells_done": ["a", "b"], "found": 12}
        stuck = repository.create_run_history(
            db_session,
            _run_data(status="running", finished_at=None, checkpoint=checkpoint),
        )
        done = repository.create_run_history(db_session, _run_data(status="success"))

        assert mark_interrupted_runs(db_session) == 1

        stuck = repository.get_run_history(db_session, stuck.id)
        assert stuck.status == "interrupted"
        assert stuck.finished_at is not None
        assert stuck.error_message
        assert stuck.checkpoint == checkpoint
        assert repository.get_run_history(db_session, done.id).status == "success"
        assert mark_interrupted_runs(db_session) == 0


class TestRunHistoryEndpoints:
    def test_list_empty(self, api_client):
        client, _ = api_client
        response = client.get("/history")
        assert response.status_code == 200
        data = response.json()
        assert data["total"] == 0
        assert data["items"] == []

    def test_list_pagination(self, api_client):
        client, seed = api_client
        seed(_run_data(location="City 0"), _run_data(location="City 1"), _run_data(location="City 2"))
        response = client.get("/history", params={"limit": 2, "offset": 0})
        data = response.json()
        assert data["total"] == 3
        assert len(data["items"]) == 2
        assert data["items"][0]["location"] == "City 2"

    def test_list_filters(self, api_client):
        client, seed = api_client
        seed(
            _run_data(run_type="search", status="success", location="NYC"),
            _run_data(run_type="agent_search", status="failed"),
        )
        response = client.get("/history", params={"run_type": "agent_search"})
        assert response.json()["total"] == 1
        response = client.get("/history", params={"status": "success"})
        assert response.json()["total"] == 1
        response = client.get("/history", params={"search": "nyc"})
        assert response.json()["total"] == 1

    def test_get_run(self, api_client):
        client, seed = api_client
        run = seed(_run_data())[0]
        response = client.get(f"/history/{run.id}")
        assert response.status_code == 200
        assert response.json()["run_type"] == "search"

    def test_get_run_404(self, api_client):
        client, _ = api_client
        response = client.get("/history/9999")
        assert response.status_code == 404
        assert response.json()["detail"]["error"] == "NotFound"

    def test_delete_run(self, api_client):
        client, seed = api_client
        run = seed(_run_data())[0]
        response = client.delete(f"/history/{run.id}")
        assert response.status_code == 204
        assert client.get(f"/history/{run.id}").status_code == 404

    def test_delete_run_404(self, api_client):
        client, _ = api_client
        response = client.delete("/history/9999")
        assert response.status_code == 404

    def test_clear_all(self, api_client):
        client, seed = api_client
        seed(_run_data(), _run_data())
        response = client.delete("/history")
        assert response.status_code == 200
        assert response.json()["deleted"] == 2
        assert client.get("/history").json()["total"] == 0

    def test_stats_endpoint(self, api_client):
        client, seed = api_client
        seed(_run_data(status="success", processing_time=1.0), _run_data(status="failed", processing_time=2.0))
        response = client.get("/history/stats")
        assert response.status_code == 200
        data = response.json()
        assert data["total_runs"] == 2
        assert data["successful"] == 1
        assert data["failed"] == 1
        assert data["avg_duration"] == 1.5

    def test_export_csv(self, api_client):
        client, seed = api_client
        seed(_run_data(location="NYC"))
        response = client.get("/history/export/csv")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/csv")
        assert "attachment; filename=run_history_" in response.headers["content-disposition"]
        rows = list(csv.DictReader(io.StringIO(response.text)))
        assert len(rows) == 1
        assert rows[0]["run_type"] == "search"
        assert rows[0]["location"] == "NYC"
        assert rows[0]["status"] == "success"

    def test_export_excel(self, api_client):
        from openpyxl import load_workbook

        client, seed = api_client
        seed(_run_data())
        response = client.get("/history/export/excel")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers["content-type"]
        workbook = load_workbook(io.BytesIO(response.content))
        worksheet = workbook.active
        assert worksheet.title == "Run History"
        headers = [cell.value for cell in worksheet[1]]
        assert "run_type" in headers
        assert "error_message" in headers
        assert worksheet.max_row == 2


class TestBusinessEndpointsLogRuns:
    @staticmethod
    def _app(tmp_path):
        from fastapi import FastAPI

        from app.api.routes import business
        from app.db.database import get_db
        from app.dependencies import get_business_search_service

        engine = create_engine(
            f"sqlite:///{tmp_path / 'bus.db'}",
            connect_args={"check_same_thread": False},
        )
        Base.metadata.create_all(bind=engine)
        SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)

        def override_get_db():
            session = SessionLocal()
            try:
                yield session
            finally:
                session.close()

        app = FastAPI()
        app.include_router(business.router)
        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[get_business_search_service] = (
            lambda: FakeBusinessService()
        )
        return app, SessionLocal, engine

    def test_search_logs_successful_run(self, tmp_path):
        from fastapi.testclient import TestClient

        from app.db import repository

        app, SessionLocal, engine = self._app(tmp_path)
        with TestClient(app) as client:
            response = client.post(
                "/search",
                json={"location": "NYC", "business_category": "cafe", "radius": 5000},
            )
            assert response.status_code == 200
            session = SessionLocal()
            try:
                runs = repository.list_run_history(session)
                assert len(runs) == 1
                assert runs[0].run_type == "search"
                assert runs[0].status == "success"
                assert runs[0].location == "NYC"
                assert runs[0].business_category == "cafe"
                assert runs[0].finished_at is not None
                assert runs[0].processing_time >= 0
            finally:
                session.close()
                engine.dispose()

    def test_agent_search_logs_run(self, tmp_path):
        from fastapi.testclient import TestClient

        from app.db import repository

        app, SessionLocal, engine = self._app(tmp_path)
        with patch("app.api.routes.business.GeminiDecisionEngine", DummyEngine), patch(
            "app.api.routes.business.LeadAgent", FakeLeadAgent
        ):
            with TestClient(app) as client:
                response = client.post(
                    "/agent/search",
                    json={
                        "location": "Austin, TX",
                        "business_category": "coffee shops",
                        "radius": 5000,
                    },
                )
                assert response.status_code == 200
                assert response.json()["total_results"] == 1
                session = SessionLocal()
                try:
                    runs = repository.list_run_history(session)
                    assert len(runs) == 1
                    assert runs[0].run_type == "agent_search"
                    assert runs[0].status == "success"
                    assert runs[0].total_results == 1
                finally:
                    session.close()
                    engine.dispose()

    def test_discover_more_logs_failed_run(self, tmp_path):
        from fastapi.testclient import TestClient

        from app.db import repository

        app, SessionLocal, engine = self._app(tmp_path)
        with TestClient(app) as client:
            response = client.post(
                "/enrich/discover", json={"fsq_place_id": "no_such_id"}
            )
            assert response.status_code == 404
            session = SessionLocal()
            try:
                runs = repository.list_run_history(session)
                assert len(runs) == 1
                assert runs[0].run_type == "discover_more"
                assert runs[0].status == "failed"
                assert runs[0].location == "no_such_id"
                assert runs[0].error_message
            finally:
                session.close()
                engine.dispose()


class FakeBusinessService:
    """Fake search service exposing the async methods the routes call."""

    async def search(self, **kwargs):
        return []

    async def get_business_details(self, fsq_place_id):
        from app.core.exceptions import NoResultsError

        raise NoResultsError(provider="foursquare", query=fsq_place_id)


class DummyEngine:
    def __init__(self, *args, **kwargs):
        pass


class FakeLeadAgent:
    def __init__(self, *args, **kwargs):
        pass

    async def search(self, **kwargs):
        return {
            "leads": [
                {
                    "business_name": "Acme Cafe",
                    "address": "1 Main St",
                    "phone": None,
                    "website": "https://acme.example.com",
                    "email": None,
                    "linkedin": None,
                    "facebook": None,
                    "instagram": None,
                    "twitter": None,
                    "youtube": None,
                    "latitude": None,
                    "longitude": None,
                    "confidence_score": 0.5,
                    "source": "foursquare",
                }
            ],
            "total_results": 1,
            "search_location": "Austin, TX",
            "search_category": "coffee shops",
            "search_radius": 5000,
            "processing_time_ms": 1234,
        }


class TestRunHistoryPageUI:
    @staticmethod
    def _load_frontend_app():
        import importlib.util

        if "frontend_app" in sys.modules:
            return sys.modules["frontend_app"]
        sys.path.insert(0, str(FRONTEND_DIR))
        try:
            spec = importlib.util.spec_from_file_location(
                "frontend_app", FRONTEND_DIR / "app.py"
            )
            module = importlib.util.module_from_spec(spec)
            sys.modules["frontend_app"] = module
            spec.loader.exec_module(module)
        finally:
            sys.path.remove(str(FRONTEND_DIR))
        return module

    @staticmethod
    def _fixture_run(run_id=1, **overrides):
        data = {
            "id": run_id,
            "run_type": "search",
            "location": "New York, NY",
            "business_category": "restaurant",
            "status": "success",
            "total_tasks": 1,
            "total_results": 5,
            "processing_time": 1.25,
            "started_at": "2026-01-01T10:00:00+00:00",
            "finished_at": "2026-01-01T10:00:01+00:00",
            "error_message": None,
        }
        data.update(overrides)
        return data

    @staticmethod
    def _fixture_stats(**overrides):
        data = {
            "total_runs": 1,
            "successful": 1,
            "failed": 0,
            "avg_duration": 1.25,
            "by_status": {"success": 1},
        }
        data.update(overrides)
        return data

    @staticmethod
    def _page(**overrides):
        from contextlib import contextmanager

        from streamlit.testing.v1 import AppTest

        harness = str(Path(__file__).parent / "_history_harness.py")

        @contextmanager
        def _run():
            appmod = TestRunHistoryPageUI._load_frontend_app()
            with patch.multiple(appmod, **overrides):
                at = AppTest.from_file(harness, default_timeout=20)
                at.run()
                yield at

        return _run()

    def test_empty_state(self):
        with self._page(
            get_run_history_stats=lambda: self._fixture_stats(),
            list_run_history=lambda **kwargs: {
                "items": [],
                "total": 0,
                "limit": kwargs.get("limit", 25),
                "offset": 0,
            },
            export_history_csv=lambda **kwargs: ("run_history.csv", b"data"),
            export_history_excel=lambda **kwargs: ("run_history.xlsx", b"data"),
        ) as at:
            assert not at.exception
            assert at.info
            assert at.button(key="history_refresh")
            assert len(at.get("download_button")) == 2

    def test_renders_table_and_summary(self):
        with self._page(
            get_run_history_stats=lambda: self._fixture_stats(),
            list_run_history=lambda **kwargs: {
                "items": [self._fixture_run()],
                "total": 1,
                "limit": kwargs.get("limit", 25),
                "offset": 0,
            },
            export_history_csv=lambda **kwargs: ("run_history.csv", b"data"),
            export_history_excel=lambda **kwargs: ("run_history.xlsx", b"data"),
        ) as at:
            assert not at.exception
            assert at.dataframe
            rendered = " ".join(str(md.value) for md in at.markdown)
            assert "run(s) found." in rendered
            assert at.button(key="history_view")
            assert at.button(key="history_delete")
            assert len(at.get("download_button")) == 2

    def test_filters_passed(self):
        calls = {}

        def fake_list(search="", run_type="", status="", limit=25, offset=0):
            calls["search"] = search
            calls["run_type"] = run_type
            calls["status"] = status
            return {"items": [], "total": 0, "limit": limit, "offset": offset}

        with self._page(
            get_run_history_stats=lambda: self._fixture_stats(),
            list_run_history=fake_list,
            export_history_csv=lambda **kwargs: ("run_history.csv", b"data"),
            export_history_excel=lambda **kwargs: ("run_history.xlsx", b"data"),
        ) as at:
            at.text_input(key="history_search").set_value("pizza")
            at.selectbox(key="history_type").set_value("agent_search")
            at.selectbox(key="history_status").set_value("failed")
            at.run()
            assert not at.exception
        assert calls["search"] == "pizza"
        assert calls["run_type"] == "agent_search"
        assert calls["status"] == "failed"

    def test_delete_flow(self):
        calls = {}

        def fake_delete(run_id):
            calls["run_id"] = run_id

        with self._page(
            get_run_history_stats=lambda: self._fixture_stats(),
            list_run_history=lambda **kwargs: {
                "items": [self._fixture_run()],
                "total": 1,
                "limit": kwargs.get("limit", 25),
                "offset": 0,
            },
            get_run_history=lambda run_id: self._fixture_run(),
            delete_run_history=fake_delete,
            export_history_csv=lambda **kwargs: ("run_history.csv", b"data"),
            export_history_excel=lambda **kwargs: ("run_history.xlsx", b"data"),
        ) as at:
            at.button(key="history_delete").click().run()
            assert not at.exception
        assert calls["run_id"] == 1

    def test_pagination_next(self):
        offsets = []

        def fake_list(search="", run_type="", status="", limit=25, offset=0):
            offsets.append(offset)
            return {
                "items": [self._fixture_run(run_id=i) for i in range(1, 6)],
                "total": 30,
                "limit": limit,
                "offset": offset,
            }

        with self._page(
            get_run_history_stats=lambda: self._fixture_stats(),
            list_run_history=fake_list,
            export_history_csv=lambda **kwargs: ("run_history.csv", b"data"),
            export_history_excel=lambda **kwargs: ("run_history.xlsx", b"data"),
        ) as at:
            assert at.button(key="history_next")
            at.button(key="history_next").click().run()
            assert not at.exception
        assert offsets[-1] == 25
