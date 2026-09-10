import csv
import io
from datetime import UTC, datetime
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.db import repository
from app.db.database import get_db
from app.db.models import RunHistory
from app.monitoring.metrics import get_metrics
from app.schemas.history import (
    RunHistoryListResponse,
    RunHistoryRead,
    RunHistoryStatsResponse,
)

router = APIRouter(prefix="/history", tags=["history"])

DbSession = Annotated[Session, Depends(get_db)]

_HISTORY_EXPORT_FIELDS = [
    "id",
    "run_type",
    "location",
    "business_category",
    "status",
    "total_tasks",
    "total_results",
    "processing_time",
    "started_at",
    "finished_at",
    "error_message",
]


def _raise_not_found(run_id: int) -> None:
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail={
            "error": "NotFound",
            "message": f"Run history entry {run_id} not found",
            "status_code": 404,
        },
    )


def _to_schema(run: RunHistory) -> RunHistoryRead:
    return RunHistoryRead.model_validate(run)


def _datetime_header() -> str:
    return datetime.now(UTC).strftime("%Y%m%d_%H%M%S")


@router.get("/stats", response_model=RunHistoryStatsResponse)
def run_history_stats(db: DbSession) -> RunHistoryStatsResponse:
    return RunHistoryStatsResponse(**repository.get_run_history_stats(db))


@router.get("", response_model=RunHistoryListResponse)
def list_run_history(
    db: DbSession,
    search: str | None = Query(default=None, max_length=200),
    run_type: str | None = Query(default=None, max_length=50),
    status_filter: str | None = Query(default=None, max_length=50, alias="status"),
    limit: int = Query(default=25, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
) -> RunHistoryListResponse:
    items = [
        _to_schema(run)
        for run in repository.list_run_history(
            db,
            search=search,
            run_type=run_type,
            status=status_filter,
            limit=limit,
            offset=offset,
        )
    ]
    total = repository.count_run_history(
        db, search=search, run_type=run_type, status=status_filter
    )
    return RunHistoryListResponse(items=items, total=total, limit=limit, offset=offset)


def _rows(
    db: Session, search: str | None, run_type: str | None, status_filter: str | None
) -> list[dict]:
    runs = repository.list_run_history(
        db,
        search=search,
        run_type=run_type,
        status=status_filter,
        limit=100000,
        offset=0,
    )
    rows = []
    for run in runs:
        row = {field: getattr(run, field, None) for field in _HISTORY_EXPORT_FIELDS}
        if row.get("started_at"):
            row["started_at"] = row["started_at"].isoformat()
        if row.get("finished_at"):
            row["finished_at"] = row["finished_at"].isoformat()
        rows.append(row)
    return rows


@router.get("/export/csv", response_class=Response)
def export_history_csv(
    db: DbSession,
    search: str | None = Query(default=None, max_length=200),
    run_type: str | None = Query(default=None, max_length=50),
    status: str | None = Query(default=None, max_length=50),
) -> Response:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=_HISTORY_EXPORT_FIELDS)
    writer.writeheader()
    for row in _rows(db, search, run_type, status):
        writer.writerow(row)

    filename = f"run_history_{_datetime_header()}.csv"
    get_metrics().increment("history_exported_csv")
    return Response(
        content=buffer.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/excel", response_class=Response)
def export_history_excel(
    db: DbSession,
    search: str | None = Query(default=None, max_length=200),
    run_type: str | None = Query(default=None, max_length=50),
    status: str | None = Query(default=None, max_length=50),
) -> Response:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Run History"
    worksheet.append(_HISTORY_EXPORT_FIELDS)
    for row in _rows(db, search, run_type, status):
        worksheet.append([row[field] for field in _HISTORY_EXPORT_FIELDS])

    for column_index in range(1, len(_HISTORY_EXPORT_FIELDS) + 1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = 22

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    filename = f"run_history_{_datetime_header()}.xlsx"
    get_metrics().increment("history_exported_excel")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/{run_id}", response_model=RunHistoryRead)
def get_run_history(run_id: int, db: DbSession) -> RunHistoryRead:
    run = repository.get_run_history(db, run_id)
    if run is None:
        _raise_not_found(run_id)
    return _to_schema(run)


@router.delete("", status_code=status.HTTP_200_OK)
def clear_run_history(db: DbSession) -> dict:
    deleted = repository.clear_run_history(db)
    return {"deleted": deleted}


@router.delete("/{run_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_run_history(run_id: int, db: DbSession) -> None:
    run = repository.get_run_history(db, run_id)
    if run is None:
        _raise_not_found(run_id)
    repository.delete_run_history(db, run)
